"""Operational reports: on screen (JSON) and as Excel / CSV downloads.

All times are converted from UTC to the hotel timezone (HOTEL_TIMEZONE, e.g. Asia/Singapore)
and the timezone is stated in every export.
"""
import csv
import io
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.orm import Session

from core.config import settings
from core.exceptions import NotFoundError, ValidationFailedError
from models.daily_sheet import DailySheet, DailySheetRoom
from models.enums import DailyRoomStatus, ItemResult, RepairStatus, RequestStatus, Role, SessionOutcome
from models.guest_request import GuestRequest
from models.repair import RepairRequest
from models.user import User

REPORTS = {
    "daily-runsheet": {"title": "Daily runsheet", "roles": {Role.ADMIN, Role.HK_SUPERVISOR}, "single_date": True},
    "ra-productivity": {"title": "Room attendant productivity", "roles": {Role.ADMIN, Role.HK_SUPERVISOR}},
    "engineering-issues": {"title": "Engineering issues", "roles": {Role.ADMIN, Role.ENG_SUPERVISOR}},
    "guest-requests": {"title": "Guest requests", "roles": {Role.ADMIN, Role.HK_SUPERVISOR}},
}

PENDING_REASON = {"NEEDS_PARTS": "Needs parts", "NEEDS_VENDOR": "Needs an outside vendor", "NEEDS_VACANT_ROOM": "Needs the room empty",
                  "NEEDS_MORE_TIME": "Needs more time", "OTHER": "Other"}


def tz() -> ZoneInfo:
    return ZoneInfo(settings.HOTEL_TIMEZONE)


def tz_label() -> str:
    abbrev = datetime.now(tz()).strftime("%Z")
    return f"{settings.HOTEL_TIMEZONE} ({abbrev})" if abbrev and not abbrev.startswith(("+", "-")) else settings.HOTEL_TIMEZONE


def local(dt: Optional[datetime]) -> Optional[str]:
    """UTC (naive) -> 'YYYY-MM-DD HH:MM' in the hotel timezone."""
    if dt is None:
        return None
    return dt.replace(tzinfo=timezone.utc).astimezone(tz()).strftime("%Y-%m-%d %H:%M")


def utc_bounds(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    """Local calendar days -> naive UTC range used in queries."""
    start = datetime.combine(date_from, time.min, tz()).astimezone(timezone.utc).replace(tzinfo=None)
    end = datetime.combine(date_to + timedelta(days=1), time.min, tz()).astimezone(timezone.utc).replace(tzinfo=None)
    return start, end


def minutes(a: Optional[datetime], b: Optional[datetime]) -> Optional[float]:
    return round((b - a).total_seconds() / 60, 1) if a and b else None


def avg(values) -> Optional[float]:
    v = [x for x in values if x is not None]
    return round(sum(v) / len(v), 1) if v else None


def _range(date_from: Optional[date], date_to: Optional[date]) -> tuple[date, date]:
    from services.daily_sheet_service import hotel_today
    date_to = date_to or hotel_today()
    date_from = date_from or (date_to - timedelta(days=6))
    if date_from > date_to:
        raise ValidationFailedError("The start date is after the end date.")
    if (date_to - date_from).days > 366:
        raise ValidationFailedError("Choose a range of one year or less.")
    return date_from, date_to


# --------------------------------------------------------------------------- reports

def daily_runsheet(db: Session, day: date) -> dict:
    sheet = db.scalar(select(DailySheet).where(DailySheet.task_date == day))
    if sheet is None:
        raise NotFoundError("No task sheet has been uploaded for this date.", code="NO_TASK_SHEET")
    from services.daily_sheet_service import current_session, priority_key
    rows = []
    for e in sorted(sheet.rooms, key=lambda e: ((e.assigned_ra.display_name if e.assigned_ra else "~"), priority_key(e))):
        s = current_session(e)
        failed = [r.item_name for r in s.results if r.result == ItemResult.FAIL] if s else []
        rows.append({
            "room": e.room_number, "sheet": e.sheet_no, "type": e.room_type_code, "fo_status": e.fo_status,
            "reservation": e.reservation_status, "arrival_today": "Yes" if e.arrival_today else "",
            "room_attendant": e.assigned_ra.display_name if e.assigned_ra else "Unassigned",
            "status": {"PENDING": "Not started", "IN_PROGRESS": "In progress", "DONE": "Done"}[e.status.value],
            "assigned_at": local(e.assigned_at), "started_at": local(e.started_at), "finished_at": local(e.finished_at),
            "minutes": minutes(e.started_at, e.finished_at),
            "result": "Issues sent to engineering" if e.outcome == SessionOutcome.FAIL else ("All passed" if e.outcome == SessionOutcome.PASS else ""),
            "failed_items": ", ".join(failed),
        })
    done = [r for r in rows if r["status"] == "Done"]
    return {
        "columns": [("room", "Room"), ("sheet", "Sheet"), ("type", "Type"), ("fo_status", "FO status"), ("reservation", "Reservation"),
                    ("arrival_today", "Arrival today"), ("room_attendant", "Room attendant"), ("status", "Status"),
                    ("assigned_at", "Assigned"), ("started_at", "Started"), ("finished_at", "Finished"), ("minutes", "Minutes"),
                    ("result", "Result"), ("failed_items", "Failed items")],
        "rows": rows,
        "summary": [("Rooms on sheet", len(rows)), ("Done", len(done)), ("Not started", sum(1 for r in rows if r["status"] == "Not started")),
                    ("Unassigned", sum(1 for r in rows if r["room_attendant"] == "Unassigned")),
                    ("Sent to engineering", sum(1 for r in rows if r["result"].startswith("Issues"))),
                    ("Average minutes per room", avg(r["minutes"] for r in done))],
        "period": day.strftime("%d %b %Y"),
    }


def ra_productivity(db: Session, date_from: date, date_to: date) -> dict:
    entries = db.scalars(select(DailySheetRoom).join(DailySheet, DailySheet.id == DailySheetRoom.sheet_id)
                         .where(DailySheet.task_date >= date_from, DailySheet.task_date <= date_to,
                                DailySheetRoom.assigned_ra_id.is_not(None))).unique().all()
    start, end = utc_bounds(date_from, date_to)
    reqs = db.scalars(select(GuestRequest).where(GuestRequest.assigned_ra_id.is_not(None),
                                                 GuestRequest.created_at >= start, GuestRequest.created_at < end)).unique().all()
    stats = defaultdict(lambda: {"assigned": 0, "done": 0, "mins": [], "issues": 0, "arrivals": 0, "req": 0, "req_done": 0, "req_mins": [], "days": set()})
    names = {}
    for e in entries:
        k = e.assigned_ra_id
        names[k] = e.assigned_ra.display_name
        st = stats[k]
        st["assigned"] += 1
        st["days"].add(e.sheet.task_date)
        st["arrivals"] += 1 if e.arrival_today else 0
        if e.status == DailyRoomStatus.DONE:
            st["done"] += 1
            st["mins"].append(minutes(e.started_at, e.finished_at))
            st["issues"] += 1 if e.outcome == SessionOutcome.FAIL else 0
    for r in reqs:
        k = r.assigned_ra_id
        names[k] = r.assigned_ra.display_name
        st = stats[k]
        st["req"] += 1
        if r.status == RequestStatus.DONE:
            st["req_done"] += 1
            st["req_mins"].append(minutes(r.assigned_at or r.created_at, r.completed_at))
    rows = [{
        "room_attendant": names[k], "days_worked": len(st["days"]), "rooms_assigned": st["assigned"], "rooms_done": st["done"],
        "completion_pct": round(100 * st["done"] / st["assigned"]) if st["assigned"] else None,
        "avg_minutes_per_room": avg(st["mins"]), "arrival_rooms": st["arrivals"], "rooms_with_issues": st["issues"],
        "requests_assigned": st["req"], "requests_done": st["req_done"], "avg_request_minutes": avg(st["req_mins"]),
    } for k, st in stats.items()]
    rows.sort(key=lambda r: (-r["rooms_done"], r["room_attendant"]))
    return {
        "columns": [("room_attendant", "Room attendant"), ("days_worked", "Days"), ("rooms_assigned", "Rooms assigned"),
                    ("rooms_done", "Rooms done"), ("completion_pct", "Done %"), ("avg_minutes_per_room", "Avg minutes per room"),
                    ("arrival_rooms", "Arrival rooms"), ("rooms_with_issues", "Rooms with issues"),
                    ("requests_assigned", "Requests assigned"), ("requests_done", "Requests done"), ("avg_request_minutes", "Avg request minutes")],
        "rows": rows,
        "summary": [("Room attendants", len(rows)), ("Rooms done", sum(r["rooms_done"] for r in rows)),
                    ("Average minutes per room", avg(m for st in stats.values() for m in st["mins"])),
                    ("Requests done", sum(r["requests_done"] for r in rows))],
    }


def engineering_issues(db: Session, date_from: date, date_to: date) -> dict:
    from services.daily_sheet_service import hotel_today
    start, end = utc_bounds(date_from, date_to)
    repairs = db.scalars(select(RepairRequest).where(RepairRequest.created_at >= start, RepairRequest.created_at < end)
                         .order_by(RepairRequest.created_at)).unique().all()
    today = hotel_today()
    now = datetime.utcnow()
    rows = []
    for r in repairs:
        pending = r.status == RepairStatus.PENDING
        days_pending = max(0, (now - r.pending_since).days) if pending and r.pending_since else None
        rows.append({
            "repair": r.repair_number, "room": r.room.room_number, "raised_by": "Front desk" if r.origin == "FRONT_DESK" else "Housekeeping",
            "problem": r.reported_defect, "priority": r.priority.value.title(), "status": r.status.value.replace("_", " ").title(),
            "technician": r.assigned_engineer.display_name if r.assigned_engineer else "", "raised_at": local(r.created_at),
            "started_at": local(r.started_at), "completed_at": local(r.completed_at),
            "minutes_to_fix": minutes(r.created_at, r.completed_at) if r.status == RepairStatus.COMPLETED else None,
            "pending_reason": PENDING_REASON.get(r.pending_reason.value, "") if pending and r.pending_reason else "",
            "pending_since": local(r.pending_since) if pending else None, "days_pending": days_pending,
            "expected_by": r.expected_date.isoformat() if pending and r.expected_date else None,
            "overdue": "Yes" if pending and r.expected_date and r.expected_date < today else "",
            "room_sellable": "" if not pending else ("Yes" if r.room_sellable else "No"),
            "escalated": "Yes" if r.escalated else "", "reopened": r.reopen_count or "",
            "diagnosis": r.diagnosis or "", "work_done": r.repair_notes or "",
        })
    pend = [x for x in rows if x["days_pending"] is not None]
    ageing = [("Pending 0-1 days", sum(1 for x in pend if x["days_pending"] <= 1)),
              ("Pending 2-3 days", sum(1 for x in pend if 2 <= x["days_pending"] <= 3)),
              ("Pending 4-7 days", sum(1 for x in pend if 4 <= x["days_pending"] <= 7)),
              ("Pending over 7 days", sum(1 for x in pend if x["days_pending"] > 7))]
    return {
        "columns": [("repair", "Repair"), ("room", "Room"), ("raised_by", "Raised by"), ("problem", "Problem"), ("priority", "Priority"),
                    ("status", "Status"), ("technician", "Technician"), ("raised_at", "Raised"), ("started_at", "Started"),
                    ("completed_at", "Completed"), ("minutes_to_fix", "Minutes to fix"), ("pending_reason", "Pending reason"),
                    ("pending_since", "Pending since"), ("days_pending", "Days pending"), ("expected_by", "Expected by"),
                    ("overdue", "Overdue"), ("room_sellable", "Room sellable"), ("escalated", "Escalated"), ("reopened", "Reopened"),
                    ("diagnosis", "Diagnosis"), ("work_done", "Work done")],
        "rows": rows,
        "summary": [("Repairs raised", len(rows)), ("Completed", sum(1 for x in rows if x["status"] == "Completed")),
                    ("Open", sum(1 for x in rows if x["status"] in ("New", "Assigned", "In Progress", "Reopened"))),
                    ("Pending (not doable now)", len(pend)), ("Overdue", sum(1 for x in rows if x["overdue"])),
                    ("Average minutes to fix", avg(x["minutes_to_fix"] for x in rows))] + ageing,
    }


def guest_requests(db: Session, date_from: date, date_to: date) -> dict:
    start, end = utc_bounds(date_from, date_to)
    reqs = db.scalars(select(GuestRequest).where(GuestRequest.created_at >= start, GuestRequest.created_at < end)
                      .order_by(GuestRequest.created_at)).unique().all()
    rows = [{
        "request": r.request_number, "room": r.room.room_number, "request_text": r.description, "priority": r.priority.value.title(),
        "status": r.status.value.replace("_", " ").title(), "raised_by": r.created_by.display_name if r.created_by else "",
        "room_attendant": r.assigned_ra.display_name if r.assigned_ra else "", "raised_at": local(r.created_at),
        "assigned_at": local(r.assigned_at), "started_at": local(r.started_at), "completed_at": local(r.completed_at),
        "minutes_to_assign": minutes(r.created_at, r.assigned_at), "minutes_to_complete": minutes(r.created_at, r.completed_at),
        "note": r.completion_note or r.cancel_reason or "",
    } for r in reqs]
    return {
        "columns": [("request", "Request"), ("room", "Room"), ("request_text", "Request text"), ("priority", "Priority"), ("status", "Status"),
                    ("raised_by", "Raised by"), ("room_attendant", "Room attendant"), ("raised_at", "Raised"), ("assigned_at", "Assigned"),
                    ("started_at", "Started"), ("completed_at", "Completed"), ("minutes_to_assign", "Minutes to assign"),
                    ("minutes_to_complete", "Minutes to complete"), ("note", "Note")],
        "rows": rows,
        "summary": [("Requests", len(rows)), ("Done", sum(1 for x in rows if x["status"] == "Done")),
                    ("Open", sum(1 for x in rows if x["status"] in ("New", "Assigned", "In Progress"))),
                    ("Cancelled", sum(1 for x in rows if x["status"] == "Cancelled")),
                    ("Average minutes to assign", avg(x["minutes_to_assign"] for x in rows)),
                    ("Average minutes to complete", avg(x["minutes_to_complete"] for x in rows))],
    }


def build(db: Session, user: User, report: str, date_from: Optional[date], date_to: Optional[date], day: Optional[date]) -> dict:
    from core.exceptions import ForbiddenError
    from services.daily_sheet_service import hotel_today
    meta = REPORTS.get(report)
    if meta is None:
        raise NotFoundError("Unknown report.")
    if user.role not in meta["roles"]:
        raise ForbiddenError("Your role does not have access to this report.")
    if meta.get("single_date"):
        the_day = day or date_to or hotel_today()
        data = daily_runsheet(db, the_day)
        file_period = the_day.isoformat()
    else:
        f, t = _range(date_from, date_to)
        data = {"daily-runsheet": None, "ra-productivity": ra_productivity, "engineering-issues": engineering_issues,
                "guest-requests": guest_requests}[report](db, f, t)
        data["period"] = f"{f:%d %b %Y} to {t:%d %b %Y}"
        file_period = f"{f.isoformat()}_to_{t.isoformat()}"
    return {
        "report": report, "title": meta["title"], "period": data["period"], "timezone": tz_label(),
        "generated_at": datetime.now(tz()).strftime("%Y-%m-%d %H:%M"),
        "columns": [{"key": k, "label": label} for k, label in data["columns"]],
        "rows": data["rows"], "summary": [{"label": label, "value": v} for label, v in data["summary"]],
        "filename": f"{report}_{file_period}",
    }


# --------------------------------------------------------------------------- files

def to_csv(rep: dict) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([f"{rep['title']} | {rep['period']} | times in {rep['timezone']} | generated {rep['generated_at']}"])
    w.writerow([c["label"] for c in rep["columns"]])
    for row in rep["rows"]:
        w.writerow(["" if row.get(c["key"]) is None else row.get(c["key"]) for c in rep["columns"]])
    return ("\ufeff" + out.getvalue()).encode("utf-8")  # BOM so Excel opens it as UTF-8


def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rep["title"][:31]
    ws["A1"] = rep["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{rep['period']}  |  times in {rep['timezone']}  |  generated {rep['generated_at']}"
    ws["A2"].font = Font(italic=True, color="555555")
    header_row = 4
    fill = PatternFill("solid", fgColor="173B57")
    for i, c in enumerate(rep["columns"], start=1):
        cell = ws.cell(row=header_row, column=i, value=c["label"])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [len(c["label"]) for c in rep["columns"]]
    for r_i, row in enumerate(rep["rows"], start=header_row + 1):
        for c_i, c in enumerate(rep["columns"], start=1):
            v = row.get(c["key"])
            ws.cell(row=r_i, column=c_i, value=v)
            widths[c_i - 1] = max(widths[c_i - 1], min(len(str(v)) if v is not None else 0, 60))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(9, w + 2)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rep["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(rep['columns']))}{header_row + len(rep['rows'])}"
    summary = wb.create_sheet("Summary")
    summary["A1"] = f"{rep['title']} summary"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"{rep['period']}  |  times in {rep['timezone']}"
    for i, s in enumerate(rep["summary"], start=4):
        summary.cell(row=i, column=1, value=s["label"]).font = Font(bold=True)
        summary.cell(row=i, column=2, value=s["value"])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
