"""Stage 4: operational reports and Excel / CSV exports, with times in the hotel timezone."""
import io
from datetime import datetime, timedelta, timezone

import pytest

from tests.conftest import engineering_items
from tests.test_daily_sheet import PNG, assign, fill, today  # noqa: F401
from tests.test_daily_sheet import hk  # noqa: F401  (fixture)


@pytest.fixture
def singapore(monkeypatch):
    from core.config import settings
    monkeypatch.setattr(settings, "HOTEL_TIMEZONE", "Asia/Singapore")


@pytest.fixture
def day_of_work(client, hk):
    """Anita finishes 105 (pass) and 202 (fail); a guest request is done; a repair goes pending."""
    assign(client, hk, ["105", "202", "108"])
    for n, fail in (("105", None), ("202", "Air conditioning")):
        eid = hk["entry"][n]["id"]
        client.post(f"/api/v1/housekeeping/runsheet/{eid}/start", headers=hk["ra1"])
        r = client.post(f"/api/v1/housekeeping/runsheet/{eid}/finish", headers=hk["ra1"], json={"items": fill(client, hk["ra1"], eid, fail=fail)})
        assert r.status_code == 200
    rep_id = r.json()["repair_ids"][0]
    req = client.post("/api/v1/requests", headers=hk["admin"], json={"room_id": hk["entry"]["108"]["room_id"], "description": "Water bottle"}).json()
    client.post(f"/api/v1/requests/{req['id']}/assign", headers=hk["sup"], json={"ra_id": hk["ra1_id"]})
    client.post(f"/api/v1/requests/{req['id']}/done", headers=hk["ra1"], json={"note": "Delivered"})
    client.post(f"/api/v1/engineering/{rep_id}/assign", headers=hk["tech"], json={"engineer_id": hk["tech_id"]})
    client.post(f"/api/v1/engineering/{rep_id}/start", headers=hk["tech"])
    client.post(f"/api/v1/engineering/{rep_id}/pending", headers=hk["tech"],
                json={"reason": "NEEDS_PARTS", "note": "Capacitor ordered", "expected_date": (today() - timedelta(days=1)).isoformat()})
    return {"repair_id": rep_id, "request_id": req["id"]}


def get(client, headers, report, fmt="json", **params):
    q = "&".join(f"{k}={v}" for k, v in params.items())
    return client.get(f"/api/v1/reports/export/{report}?format={fmt}&{q}", headers=headers)


def test_daily_runsheet_report(client, hk, day_of_work, singapore):
    r = get(client, hk["sup"], "daily-runsheet", day=today().isoformat())
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["title"] == "Daily runsheet" and "Asia/Singapore" in rep["timezone"]
    rows = {x["room"]: x for x in rep["rows"]}
    assert len(rows) == 90
    assert rows["105"]["status"] == "Done" and rows["105"]["result"] == "All passed" and rows["105"]["room_attendant"] == "Anita"
    assert rows["202"]["result"] == "Issues sent to engineering" and rows["202"]["failed_items"] == "Air conditioning"
    assert rows["202"]["arrival_today"] == "Yes" and rows["108"]["status"] == "Not started"
    summary = {s["label"]: s["value"] for s in rep["summary"]}
    assert summary["Done"] == 2 and summary["Sent to engineering"] == 1 and summary["Unassigned"] == 87
    # finished time is shown in Singapore time (UTC+8)
    from database.connection import SessionLocal
    from models.daily_sheet import DailySheetRoom
    with SessionLocal() as db:
        finished_utc = db.get(DailySheetRoom, hk["entry"]["105"]["id"]).finished_at
    assert rows["105"]["finished_at"] == (finished_utc + timedelta(hours=8)).strftime("%Y-%m-%d %H:%M")
    assert get(client, hk["sup"], "daily-runsheet", day="2020-01-01").status_code == 404


def test_ra_productivity(client, hk, day_of_work):
    rep = get(client, hk["sup"], "ra-productivity", date_from=today().isoformat(), date_to=today().isoformat()).json()
    anita = next(r for r in rep["rows"] if r["room_attendant"] == "Anita")
    assert anita["rooms_assigned"] == 3 and anita["rooms_done"] == 2 and anita["completion_pct"] == 67
    assert anita["rooms_with_issues"] == 1 and anita["arrival_rooms"] == 1
    assert anita["requests_assigned"] == 1 and anita["requests_done"] == 1 and anita["avg_minutes_per_room"] is not None


def test_engineering_issues_with_pending_ageing(client, hk, day_of_work):
    rep = get(client, hk["admin"], "engineering-issues").json()  # default: last 7 days
    row = next(r for r in rep["rows"] if r["room"] == "202")
    assert row["status"] == "Pending" and row["pending_reason"] == "Needs parts" and row["overdue"] == "Yes"
    assert row["raised_by"] == "Housekeeping" and row["technician"] == "Sunil" and row["days_pending"] == 0
    summary = {s["label"]: s["value"] for s in rep["summary"]}
    assert summary["Pending (not doable now)"] == 1 and summary["Overdue"] == 1 and summary["Pending 0-1 days"] == 1


def test_guest_requests_report(client, hk, day_of_work):
    rep = get(client, hk["sup"], "guest-requests").json()
    assert len(rep["rows"]) == 1
    row = rep["rows"][0]
    assert row["status"] == "Done" and row["room_attendant"] == "Anita" and row["note"] == "Delivered"
    assert row["minutes_to_complete"] is not None and row["raised_at"] and row["completed_at"]


def test_exports_open_in_excel_and_csv(client, hk, day_of_work, singapore):
    from openpyxl import load_workbook
    r = get(client, hk["sup"], "daily-runsheet", "xlsx", day=today().isoformat())
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f'daily-runsheet_{today().isoformat()}.xlsx' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Daily runsheet"]
    assert ws["A1"].value == "Daily runsheet" and "Asia/Singapore" in ws["A2"].value
    assert [c.value for c in ws[4]][:3] == ["Room", "Sheet", "Type"] and ws.max_row == 4 + 90
    assert "Summary" in wb.sheetnames
    r = get(client, hk["sup"], "guest-requests", "csv")
    text = r.content.decode("utf-8-sig").splitlines()
    assert r.headers["content-type"].startswith("text/csv") and "times in Asia/Singapore" in text[0]
    assert text[1].startswith("Request,Room,Request text") and "Water bottle" in text[2]


def test_report_permissions(client, hk, day_of_work):
    assert get(client, hk["sup"], "engineering-issues").status_code == 403
    assert get(client, hk["tech"], "guest-requests").status_code == 403
    assert get(client, hk["ra1"], "daily-runsheet").status_code == 403
    assert get(client, hk["admin"], "ra-productivity").status_code == 200
    from tests.conftest import supervisor
    eng_sup = supervisor(client, hk["admin"], "ENGINEERING")
    assert get(client, eng_sup, "engineering-issues").status_code == 200
    assert get(client, eng_sup, "guest-requests").status_code == 403
    assert get(client, hk["sup"], "ra-productivity", date_from="2026-10-10", date_to="2026-10-01").status_code == 422
